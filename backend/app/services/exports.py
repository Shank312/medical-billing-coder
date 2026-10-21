"""Secure export generation (CSV/XLSX/PDF) with tenant scoping and sanitization.

Exports reflect persisted database state only — they never regenerate AI
suggestions or modify review decisions. An AI suggestion is labeled approved only
when an explicit persisted review decision establishes it.
"""
import csv
import io

from sqlalchemy import select
from sqlalchemy.orm import Session

from ..core.config import get_settings
from ..core.errors import BadRequestError
from ..core.time import utcnow
from ..models.encounter import Encounter
from ..models.ingestion import IngestionDocument
from ..repositories.coding import SuggestionRepository
from ..repositories.encounters import EncounterRepository
from ..repositories.ingestion import BatchRepository
from .serializers import _parse_str_list, latest_job, review_status

_FORMULA_TRIGGERS = ("=", "+", "-", "@")

CSV_COLUMNS = [
    "encounter_id",
    "date_of_service",
    "patient_reference",
    "encounter_type",
    "place_of_service",
    "coding_status",
    "source",
    "code_system",
    "code",
    "description",
    "confidence",
    "validation_status",
    "review_decision",
    "corrected_code",
    "evidence_quotes",
    "missing_documentation",
    "contradictions",
    "reviewed_at",
]

BATCH_CSV_COLUMNS = [
    "batch_id",
    "batch_status",
    "date_of_service",
    "document_id",
    "original_filename",
    "media_type",
    "extraction_status",
    "error_code",
    "encounter_id",
    "coding_job_id",
    "created_at",
    "completed_at",
]


def sanitize_cell(value) -> str:
    """Neutralize spreadsheet formula injection; return a plain string."""
    if value is None:
        return ""
    text = str(value)
    stripped = text.lstrip("\t\r\n ")
    if stripped.startswith(_FORMULA_TRIGGERS):
        return "'" + text
    return text


def _review_decision(status: str) -> str:
    return "unreviewed" if status == "pending" else status


def _joined(items) -> str:
    return "; ".join(str(i) for i in (items or []))


def collect_export_data(
    db: Session,
    organization_id: str,
    *,
    search: str | None = None,
    status: str | None = None,
    sort: str | None = None,
    order: str | None = None,
) -> dict:
    """Collect tenant-scoped export data (bounded by ``max_export_rows``)."""
    settings = get_settings()
    encounter_repo = EncounterRepository(db, organization_id)
    suggestion_repo = SuggestionRepository(db, organization_id)

    encounters = encounter_repo.list(
        search=search,
        status=status,
        sort=sort,
        order=order,
        limit=settings.max_export_rows + 1,
        offset=0,
    )
    if len(encounters) > settings.max_export_rows:
        raise BadRequestError(
            f"Export exceeds the {settings.max_export_rows} row limit; narrow the filters",
            code="export_too_large",
        )

    rows: list[dict] = []
    encounter_rows: list[dict] = []
    review_rows: list[dict] = []
    warning_rows: list[dict] = []

    for encounter in encounters:
        job = latest_job(db, encounter.id)
        coding_status = job.status if job else None
        enc_review = review_status(db, encounter.id)
        missing = _parse_str_list(job.missing_documentation_json) if job else []
        contradictions = _parse_str_list(job.contradictions_json) if job else []

        encounter_rows.append(
            {
                "encounter_id": encounter.id,
                "date_of_service": encounter.date_of_service.isoformat(),
                "patient_reference": encounter.patient_reference,
                "encounter_type": encounter.encounter_type,
                "place_of_service": encounter.place_of_service,
                "coding_status": coding_status,
                "review_status": enc_review,
            }
        )
        for item in missing:
            warning_rows.append({"encounter_id": encounter.id, "type": "missing_documentation", "item": item})
        for item in contradictions:
            warning_rows.append({"encounter_id": encounter.id, "type": "contradiction", "item": item})

        suggestions = suggestion_repo.list_for_encounter(encounter.id)
        if not suggestions:
            rows.append(
                {
                    "encounter_id": encounter.id,
                    "date_of_service": encounter.date_of_service.isoformat(),
                    "patient_reference": encounter.patient_reference,
                    "encounter_type": encounter.encounter_type,
                    "place_of_service": encounter.place_of_service,
                    "coding_status": coding_status,
                    "source": "",
                    "code_system": "",
                    "code": "",
                    "description": "",
                    "confidence": "",
                    "validation_status": "",
                    "review_decision": "",
                    "corrected_code": "",
                    "evidence_quotes": "",
                    "missing_documentation": _joined(missing),
                    "contradictions": _joined(contradictions),
                    "reviewed_at": "",
                }
            )
            continue

        for suggestion in suggestions:
            evidence = suggestion_repo.list_evidence(suggestion.id)
            decision = _review_decision(suggestion.review_status)
            rows.append(
                {
                    "encounter_id": encounter.id,
                    "date_of_service": encounter.date_of_service.isoformat(),
                    "patient_reference": encounter.patient_reference,
                    "encounter_type": encounter.encounter_type,
                    "place_of_service": encounter.place_of_service,
                    "coding_status": coding_status,
                    "source": "ai_suggested",
                    "code_system": suggestion.code_type,
                    "code": suggestion.code,
                    "description": suggestion.description,
                    "confidence": suggestion.confidence,
                    "validation_status": suggestion.validation_status,
                    "review_decision": decision,
                    "corrected_code": suggestion.corrected_code,
                    "evidence_quotes": _joined(e.quote for e in evidence),
                    "missing_documentation": _joined(missing),
                    "contradictions": _joined(contradictions),
                    "reviewed_at": suggestion.reviewed_at.isoformat() if suggestion.reviewed_at else "",
                }
            )
            if suggestion.review_status != "pending":
                review_rows.append(
                    {
                        "encounter_id": encounter.id,
                        "code": suggestion.code,
                        "decision": decision,
                        "corrected_code": suggestion.corrected_code,
                        "reviewed_at": suggestion.reviewed_at.isoformat() if suggestion.reviewed_at else "",
                    }
                )

    return {
        "rows": rows,
        "encounters": encounter_rows,
        "reviews": review_rows,
        "warnings": warning_rows,
    }


def generate_csv(data: dict, columns: list[str]) -> str:
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=columns, extrasaction="ignore")
    writer.writeheader()
    for row in data["rows"]:
        writer.writerow({k: sanitize_cell(v) for k, v in row.items()})
    return buffer.getvalue()


def generate_xlsx(data: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)

    def add_sheet(title: str, columns: list[str], rows: list[dict], widths: dict[str, int] | None = None):
        ws = wb.create_sheet(title=title)
        ws.append(columns)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append([sanitize_cell(row.get(c, "")) for c in columns])
        ws.freeze_panes = "A2"
        for idx, col in enumerate(columns, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = (widths or {}).get(col, 18)
        return ws

    add_sheet(
        "Encounters",
        ["encounter_id", "date_of_service", "patient_reference", "encounter_type", "place_of_service", "coding_status", "review_status"],
        data["encounters"],
    )
    add_sheet(
        "Coding Results",
        ["encounter_id", "source", "code_system", "code", "description", "confidence", "validation_status", "review_decision", "corrected_code"],
        [r for r in data["rows"] if r["code"]],
    )
    add_sheet(
        "Review History",
        ["encounter_id", "code", "decision", "corrected_code", "reviewed_at"],
        data["reviews"],
    )
    add_sheet(
        "Warnings",
        ["encounter_id", "type", "item"],
        data["warnings"],
    )

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _pdf_escape(text: str) -> str:
    return str(text).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def generate_pdf_summary(db: Session, organization_id: str, encounter: Encounter) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    suggestion_repo = SuggestionRepository(db, organization_id)
    job = latest_job(db, encounter.id)
    suggestions = suggestion_repo.list_for_encounter(encounter.id)
    missing = _parse_str_list(job.missing_documentation_json) if job else []
    contradictions = _parse_str_list(job.contradictions_json) if job else []
    unreviewed = any(s.review_status == "pending" for s in suggestions)

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("TitleX", parent=styles["Title"], fontSize=18, spaceAfter=10)
    heading_style = ParagraphStyle("HeadingX", parent=styles["Heading2"], fontSize=13, spaceBefore=10, spaceAfter=4)
    body_style = styles["BodyText"]
    small_style = ParagraphStyle("Small", parent=styles["BodyText"], fontSize=9, textColor=colors.HexColor("#555555"))

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, title="Coding Summary")
    story = []

    story.append(Paragraph("Coding Summary", title_style))
    story.append(Paragraph(
        f"Encounter {_pdf_escape(encounter.id)} · Date of service: {encounter.date_of_service.isoformat()}",
        body_style,
    ))
    if encounter.patient_reference:
        story.append(Paragraph(f"Patient reference: {_pdf_escape(encounter.patient_reference)}", body_style))
    if encounter.encounter_type:
        story.append(Paragraph(f"Type: {_pdf_escape(encounter.encounter_type)}", body_style))
    if encounter.place_of_service:
        story.append(Paragraph(f"Place of service: {_pdf_escape(encounter.place_of_service)}", body_style))
    story.append(Paragraph(f"Coding status: {_pdf_escape(job.status if job else 'not coded')}", body_style))

    if unreviewed:
        story.append(Spacer(1, 8))
        story.append(Paragraph(
            "<b>UNREVIEWED</b> — this coding output has not been reviewed and requires qualified human review.",
            ParagraphStyle("Warn", parent=body_style, textColor=colors.HexColor("#9a3412"), backColor=colors.HexColor("#fef2f2"), borderPadding=6),
        ))

    story.append(Paragraph("Coding suggestions", heading_style))
    if suggestions:
        table_data = [["System", "Code", "Description", "Review", "Corrected", "Validation"]]
        for s in suggestions:
            table_data.append([
                s.code_type,
                s.code,
                (s.description or "")[:80],
                _review_decision(s.review_status),
                s.corrected_code or "",
                s.validation_status or "",
            ])
        table = Table(table_data, repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f1f5f9")),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(table)
    else:
        story.append(Paragraph("No coding suggestions.", body_style))

    if missing:
        story.append(Paragraph("Missing documentation", heading_style))
        for item in missing:
            story.append(Paragraph(f"• {_pdf_escape(item)}", body_style))
    if contradictions:
        story.append(Paragraph("Contradictions", heading_style))
        for item in contradictions:
            story.append(Paragraph(f"• {_pdf_escape(item)}", body_style))

    story.append(Spacer(1, 16))
    story.append(Paragraph(f"Generated {utcnow().isoformat()} UTC.", small_style))
    story.append(Paragraph(
        "This summary is an informational coding aid, not a submitted claim, a guarantee of "
        "reimbursement, or a certification of coding correctness.",
        small_style,
    ))

    doc.build(story)
    return buffer.getvalue()


def collect_batch_export(db: Session, organization_id: str, batch_id: str) -> dict:
    batch = BatchRepository(db, organization_id).get(batch_id)
    documents = db.execute(
        select(IngestionDocument)
        .where(IngestionDocument.batch_id == batch_id)
        .order_by(IngestionDocument.created_at)
    ).scalars().all()
    rows = [
        {
            "batch_id": batch.id,
            "batch_status": batch.status,
            "date_of_service": batch.date_of_service.isoformat() if batch.date_of_service else "",
            "document_id": d.id,
            "original_filename": d.original_filename,
            "media_type": d.media_type,
            "extraction_status": d.extraction_status,
            "error_code": d.error_code,
            "encounter_id": d.encounter_id,
            "coding_job_id": d.coding_job_id,
            "created_at": d.created_at.isoformat(),
            "completed_at": d.completed_at.isoformat() if d.completed_at else "",
        }
        for d in documents
    ]
    return {"rows": rows}
